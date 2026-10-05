#!/usr/bin/env python3
"""
Extract text from all PPTX files in docs/slides/ and generate organized raw text + summaries.

For each PPTX:
  - Creates docs/slides/<name>/ subdirectory
  - Saves raw.txt (all slide text, slide by slide)
  - Saves summary.md if raw text > 10KB

Slides are order-sensitive, so raw.txt is prefixed with slide numbers and
separators. Text is pulled from text frames, tables, grouped shapes, and
speaker notes.
"""

import os
import re
import sys
import traceback

try:
    from pptx import Presentation
    from pptx.util import Emu
except ImportError:
    print("python-pptx is required: pip install python-pptx")
    sys.exit(1)

SLIDES_DIR = os.path.join("docs", "slides")

# ---- Helper: sanitize pptx filename to directory name ----
def pptx_to_dirname(filename):
    """Convert a PPTX filename to a suitable directory name."""
    name = filename
    # Remove .pptx extension
    if name.lower().endswith(".pptx"):
        name = name[:-5]
    # Replace spaces and special chars with dashes, keep alphanumeric
    name = re.sub(r'[\s_]+', '-', name.strip())
    name = re.sub(r'[^a-zA-Z0-9\-()]', '', name)
    return name

def shape_sort_key(shape):
    """Order shapes top-to-bottom, then left-to-right."""
    try:
        top = shape.top if shape.top is not None else 0
        left = shape.left if shape.left is not None else 0
        return (top, left)
    except Exception:
        return (0, 0)

def extract_shape_text(shape, lines, depth=0):
    """Recursively pull text out of a shape (handles groups)."""
    # Grouped shapes: recurse into members
    if shape.shape_type == 6 or getattr(shape, "shapes", None) is not None and hasattr(shape, "shapes"):
        try:
            members = list(shape.shapes)
        except (AttributeError, NotImplementedError):
            members = []
        if members:
            for sub in sorted(members, key=shape_sort_key):
                extract_shape_text(sub, lines, depth + 1)
            return

    # Tables: render cell text row by row
    if getattr(shape, "has_table", False) and shape.has_table:
        for row in shape.table.rows:
            cells = [c.text.replace("\n", " ").strip() for c in row.cells]
            if any(cells):
                lines.append("[table] " + " | ".join(cells))
        return

    # Text frames
    if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
        for para in shape.text_frame.paragraphs:
            # para.text preserves soft line breaks as \x0b (vertical tab).
            # Joining run.text would drop them and glue lines together.
            text = para.text.strip()
            if not text:
                continue
            indent = "  " * depth
            for soft_line in text.split("\x0b"):
                soft_line = soft_line.strip()
                if soft_line:
                    lines.append(f"{indent}{soft_line}")

def extract_slide_text(slide):
    """Return list of text lines for one slide, in visual order."""
    lines = []
    shapes = list(slide.shapes)
    for shape in sorted(shapes, key=shape_sort_key):
        extract_shape_text(shape, lines)
    return lines

FOOTER_MARKER = "Design & Analysis of Experiments"

def slide_title(lines):
    """
    Pick a slide's title: the first content line that is not a repeated
    footer or bare slide number.
    """
    for line in lines:
        s = line.strip()
        if not s or s.startswith("[notes]") or s.startswith("[table]"):
            continue
        if FOOTER_MARKER in s or re.fullmatch(r"\d+", s):
            continue
        return s
    return "(no title text)"

def simple_summarize(text, max_chars=8000):
    """
    Generate a summary from extracted slide text:
    - Numbered slide list with each slide's title
    - Count of slides carrying speaker notes
    """
    blocks = re.split(r"^## Slide (\d+)$", text, flags=re.M)
    entries = []
    # blocks[0] is the file header; then alternating (number, body)
    for i in range(1, len(blocks) - 1, 2):
        num, body = blocks[i], blocks[i + 1]
        body_lines = body.strip().split("\n")
        entries.append((num, slide_title(body_lines), body.count("[notes]")))

    notes_slides = [n for n, _, has in entries if has]

    parts = ["## Slides\n"]
    for num, title, has in entries:
        suffix = " _(has notes)_" if has else ""
        parts.append(f"- Slide {num}: {title}{suffix}")

    parts.append("")
    parts.append("## Stats\n")
    parts.append(f"- Total slides: {len(entries)}")
    parts.append(f"- Slides with speaker notes: {len(notes_slides)}"
                 + (f" ({', '.join(notes_slides)})" if notes_slides else ""))
    parts.append(f"- Total characters of extracted text: {len(text)}")

    return "\n".join(parts)[:max_chars] + "\n"

def main():
    if not os.path.isdir(SLIDES_DIR):
        print(f"Slides directory not found: {SLIDES_DIR}")
        sys.exit(1)

    pptx_files = sorted(
        os.path.join(SLIDES_DIR, f)
        for f in os.listdir(SLIDES_DIR)
        if f.lower().endswith(".pptx")
    )

    print(f"Found {len(pptx_files)} PPTX files in {SLIDES_DIR}\n")

    extracted_count = 0
    summarized_count = 0
    skipped_count = 0
    error_count = 0

    for pptx_path in pptx_files:
        filename = os.path.basename(pptx_path)
        dirname = pptx_to_dirname(filename)
        out_dir = os.path.join(SLIDES_DIR, dirname)
        os.makedirs(out_dir, exist_ok=True)
        raw_path = os.path.join(out_dir, "raw.txt")

        # Skip if already extracted
        if os.path.exists(raw_path) and os.path.getsize(raw_path) > 0:
            print(f"  SKIP (existing): {filename} -> {raw_path}")
            skipped_count += 1
            continue

        print(f"  EXTRACTING: {filename} -> {raw_path}")

        try:
            prs = Presentation(pptx_path)
            out = [f"# Source: {filename}", f"# Slides: {len(prs.slides)}", ""]

            for i, slide in enumerate(prs.slides, start=1):
                out.append(f"## Slide {i}")
                lines = extract_slide_text(slide)
                out.extend(lines)

                # Speaker notes
                try:
                    if slide.has_notes_slide:
                        notes_text = slide.notes_slide.notes_text_frame.text.strip()
                        if notes_text:
                            out.append("[notes]")
                            for nline in notes_text.split("\n"):
                                if nline.strip():
                                    out.append(f"  {nline.strip()}")
                except (AttributeError, NotImplementedError):
                    pass

                out.append("")

            text = "\n".join(out)
            with open(raw_path, "w", encoding="utf-8", errors="replace") as f:
                f.write(text)

            file_size = os.path.getsize(raw_path)
            extracted_count += 1
            print(f"  -> {file_size} bytes, {len(prs.slides)} slides")

            if file_size > 10000:
                summary_path = os.path.join(out_dir, "summary.md")
                summary = simple_summarize(text)
                with open(summary_path, "w", encoding="utf-8", errors="replace") as f:
                    f.write(f"# Summary: {filename}\n\n")
                    f.write(summary)
                summarized_count += 1
                print(f"  -> Summary: {os.path.getsize(summary_path)} bytes")

        except Exception as e:
            print(f"  ERROR: {e}")
            traceback.print_exc()
            error_count += 1

    print(f"\n{'='*60}")
    print(f"Done: {extracted_count} extracted, {summarized_count} summarized, "
          f"{skipped_count} skipped, {error_count} errors")

if __name__ == "__main__":
    main()
